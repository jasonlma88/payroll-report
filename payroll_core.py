#!/usr/bin/env python3
"""
Payroll report engine.

Reads a timesheet export (the "Payroll Report (jobs)" sheet produced by the
time-tracking system) and collapses it into one row per employee with
Sick Leave / Vacation / Work Hours / Personal Total.

Depends only on openpyxl and the standard library — no pandas. That keeps the
frozen application small enough to hand to someone as a single download.

No UI, no printing, no reliance on the current working directory: everything is
reported back through a ProcessResult so the CLI and the GUI can render it
however they like.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Columns we expect in the timesheet export.
EXPECTED_COLUMNS = ["Person", "Description", "Project", "Hours"]

# How a project/charge code maps to a payroll bucket. First match wins; anything
# that matches nothing is Work Hours (which is what keeps "Company Holiday" and
# every client project in the Work Hours column).
CATEGORY_RULES = [
    ("Sick Leave", ("cd-sick", "sick leave")),
    ("Vacation", ("cd-vac", "paid vacation")),
]

BUCKETS = ["Sick Leave", "Vacation", "Work Hours"]
OUTPUT_COLUMNS = ["Employee"] + BUCKETS + ["Personal Total"]

# Rows in the Description column that are per-employee subtotals, not work.
SUMMARY_DESCRIPTIONS = {"worked time"}

# Tolerance when cross-checking our totals against the file's own subtotals.
TOLERANCE = 0.01

MAX_HEADER_SEARCH_ROWS = 25


class PayrollError(Exception):
    """A problem the user can act on (bad file, wrong layout, locked output)."""


@dataclass
class ProcessResult:
    """Everything a caller needs to report on a run."""

    rows: list                                # list[dict] keyed by OUTPUT_COLUMNS
    source: Path
    period: str | None = None
    header_row: int | None = None             # 1-based row number in the sheet
    sheet_name: str | None = None
    entry_counts: dict = field(default_factory=dict)
    warnings: list = field(default_factory=list)

    @property
    def employee_count(self) -> int:
        return len(self.rows)

    @property
    def total_hours(self) -> float:
        return sum(row["Personal Total"] for row in self.rows)

    def with_bucket(self, bucket: str) -> list:
        """Rows that have any hours in `bucket`, for summary display."""
        return [row for row in self.rows if row[bucket]]

    @property
    def suggested_filename(self) -> str:
        if self.period:
            return f"Payroll_{self.period}.xlsx"
        return f"Payroll_{datetime.now():%Y-%m-%d}.xlsx"


def safe_path(value) -> Path:
    """Path(value) with ~ expansion, but never for Excel's ~$lock.xlsx names."""
    path = Path(value)
    if path.name.startswith("~$") or (path.parts and path.parts[0].startswith("~$")):
        return path
    try:
        return path.expanduser()
    except RuntimeError:            # no home directory available
        return path


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------

def _cell_text(value) -> str:
    """A cell as trimmed text ('' for blanks)."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _to_hours(value) -> float | None:
    """Coerce a cell to a number of hours, or None if it isn't one."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _check_extension(path: Path) -> None:
    """Reject file types we cannot read, with advice rather than a stack trace."""
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise PayrollError(
            f"'{path.name}' is in the old .xls format, which this tool cannot read.\n\n"
            "Open it in Excel and use File > Save As to save it as .xlsx, then try again."
        )
    if suffix not in (".xlsx", ".xlsm"):
        raise PayrollError(
            f"'{path.name}' is not an Excel workbook. "
            "Expected a .xlsx timesheet export."
        )


def _read_sheet(path: Path) -> tuple[list, str]:
    """Load a sheet as a list of row-tuples, preferring one with our header."""
    _check_extension(path)

    try:
        book = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise PayrollError(
            f"Could not open '{path.name}': {exc}\n\n"
            "If the file is open in Excel, close it and try again."
        ) from exc

    try:
        sheets = {name: [row for row in book[name].iter_rows(values_only=True)]
                  for name in book.sheetnames}
    finally:
        book.close()

    if not sheets:
        raise PayrollError(f"'{path.name}' contains no sheets.")

    for name, rows in sheets.items():
        if _find_header_row(rows) is not None:
            return rows, name
    first = next(iter(sheets))
    return sheets[first], first


def _find_header_row(rows: list) -> int | None:
    """Return the 0-based index of the row holding the Person/.../Hours header."""
    wanted = {c.lower() for c in EXPECTED_COLUMNS}
    for idx in range(min(MAX_HEADER_SEARCH_ROWS, len(rows))):
        values = {_cell_text(v).lower() for v in rows[idx] if _cell_text(v)}
        if wanted.issubset(values):
            return idx
    return None


def _extract_period(rows: list, header_idx: int) -> str | None:
    """Pull the pay period ('8/9/26 - 8/22/26') out of the rows above the header."""
    pattern = re.compile(
        r"(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\s*(?:-|–|—|to)\s*"
        r"(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})"
    )
    for idx in range(header_idx):
        for value in rows[idx]:
            match = pattern.search(_cell_text(value))
            if not match:
                continue
            parts = []
            for raw in match.groups():
                for fmt in ("%m/%d/%y", "%m/%d/%Y", "%m-%d-%y", "%m-%d-%Y"):
                    try:
                        parts.append(datetime.strptime(raw, fmt).strftime("%Y-%m-%d"))
                        break
                    except ValueError:
                        continue
                else:
                    return None
            return f"{parts[0]}_to_{parts[1]}"
    return None


def _categorize(project: str) -> str:
    lowered = project.lower()
    for bucket, needles in CATEGORY_RULES:
        if any(needle in lowered for needle in needles):
            return bucket
    return "Work Hours"


# --------------------------------------------------------------------------
# Processing
# --------------------------------------------------------------------------

def process_timesheet(input_path) -> ProcessResult:
    """Turn a timesheet export into a per-employee payroll report."""
    path = safe_path(input_path)
    try:
        path = path.resolve()
    except OSError:
        pass

    # Name-based checks first: they give a specific message whether or not the
    # path exists, which is what a user pointed at the wrong file needs to hear.
    if path.name.startswith("~$"):
        raise PayrollError(
            f"'{path.name}' is an Excel lock file, not a timesheet.\n\n"
            "Pick the real workbook — the one without the ~$ at the start of its name."
        )
    _check_extension(path)
    if not path.exists():
        raise PayrollError(f"File not found: {path}")

    rows, sheet_name = _read_sheet(path)
    header_idx = _find_header_row(rows)
    if header_idx is None:
        raise PayrollError(
            f"Could not find a header row with {', '.join(EXPECTED_COLUMNS)} in "
            f"'{path.name}' (sheet '{sheet_name}').\n\n"
            "Is this the timesheet export, rather than a finished payroll report?"
        )

    period = _extract_period(rows, header_idx)

    # Map header labels to column positions so column order can move.
    header = {
        _cell_text(v).lower(): pos
        for pos, v in enumerate(rows[header_idx])
        if _cell_text(v)
    }
    cols = {name: header[name.lower()] for name in EXPECTED_COLUMNS}
    width = max(cols.values()) + 1

    totals: dict[str, dict] = {}
    order: list[str] = []
    stated: dict[str, float] = {}
    counts = {bucket: 0 for bucket in BUCKETS}
    warnings: list[str] = []
    current = None

    for offset, raw in enumerate(rows[header_idx + 1:]):
        sheet_row = header_idx + 2 + offset       # 1-based row number in the sheet
        if len(raw) < width:
            raw = tuple(raw) + (None,) * (width - len(raw))

        person = _cell_text(raw[cols["Person"]])
        description = _cell_text(raw[cols["Description"]])
        project = _cell_text(raw[cols["Project"]])
        hours = _to_hours(raw[cols["Hours"]])

        # An employee row: starts a new employee and states their own total.
        if person:
            if person.lower() == "person":        # a repeated header block
                # Clear the current employee. Leaving it set would bill the
                # rows after this header to whoever came before it.
                current = None
                continue
            current = person
            if current not in totals:
                totals[current] = {"Employee": current, **{b: 0.0 for b in BUCKETS}}
                order.append(current)
            if hours is not None:
                stated[current] = stated.get(current, 0.0) + hours
            continue

        # Per-employee subtotal rows carry no new information.
        if description.lower() in SUMMARY_DESCRIPTIONS:
            continue
        if not project:
            continue

        if current is None:
            warnings.append(
                f"Row {sheet_row}: '{project}' appears before any employee name; skipped."
            )
            continue
        if hours is None:
            warnings.append(
                f"Row {sheet_row}: {current} / '{project}' has no readable Hours value; skipped."
            )
            continue

        bucket = _categorize(project)
        totals[current][bucket] += hours
        counts[bucket] += 1

    if not totals:
        raise PayrollError(
            f"No employees found in '{path.name}'. The sheet has a header but no data rows."
        )

    report_rows = []
    for name in order:
        entry = totals[name]
        total = sum(entry[b] for b in BUCKETS)
        # Blank out zero buckets for display; keep Personal Total numeric.
        report_rows.append(
            {
                "Employee": name,
                **{b: (entry[b] if entry[b] else "") for b in BUCKETS},
                "Personal Total": total,
            }
        )
        declared = stated.get(name)
        if declared is not None and abs(declared - total) > TOLERANCE:
            warnings.append(
                f"{name}: file says {declared:g} hours, "
                f"the line items add up to {total:g}."
            )

    return ProcessResult(
        rows=report_rows,
        source=path,
        period=period,
        header_row=header_idx + 1,
        sheet_name=sheet_name,
        entry_counts=counts,
        warnings=warnings,
    )


# --------------------------------------------------------------------------
# Writing
# --------------------------------------------------------------------------

# Leading characters that Excel and Sheets treat as the start of a formula.
FORMULA_LEAD = ("=", "+", "-", "@")


def force_text(cell) -> None:
    """Keep a name like '=Smith' as text instead of letting Excel run it.

    openpyxl infers a formula from the leading '=', which both destroys the
    name on read-back and hands anyone opening the report an executable cell.
    """
    if isinstance(cell.value, str) and cell.value[:1] in FORMULA_LEAD:
        cell.data_type = "s"
        cell.quotePrefix = True


def unique_path(path: Path) -> Path:
    """Return `path`, or path with ' (2)', ' (3)'... if it is already taken."""
    if not path.exists():
        return path
    for n in range(2, 1000):
        candidate = path.with_name(f"{path.stem} ({n}){path.suffix}")
        if not candidate.exists():
            return candidate
    raise PayrollError(f"Too many existing files named like '{path.name}'.")


def save_report(result: ProcessResult, output_path, overwrite: bool = False) -> Path:
    """Write the report to Excel. Returns the path actually written."""
    path = safe_path(output_path)
    if path.is_dir() or not path.suffix:
        # No extension means the user named a folder, whether or not it exists yet.
        path = path / result.suggested_filename
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise PayrollError(f"Could not create the folder '{path.parent}': {exc}") from exc

    if not overwrite:
        path = unique_path(path)

    book = Workbook()
    sheet = book.active
    sheet.title = "Payroll Report"

    sheet.append(OUTPUT_COLUMNS)
    for row in result.rows:
        # Empty strings become truly empty cells, matching the reference report.
        sheet.append([row[c] if row[c] != "" else None for c in OUTPUT_COLUMNS])
        force_text(sheet.cell(row=sheet.max_row, column=1))

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", start_color="E6E6FA", end_color="E6E6FA")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for pos in range(1, sheet.max_column + 1):
        letter = get_column_letter(pos)
        widest = max(
            (len(str(c.value)) for c in sheet[letter] if c.value is not None), default=0
        )
        sheet.column_dimensions[letter].width = min(widest + 2, 50)

    sheet.freeze_panes = "A2"

    try:
        book.save(path)
    except PermissionError as exc:
        raise PayrollError(
            f"Cannot write '{path.name}' — it is probably open in Excel.\n\n"
            "Close it and try again."
        ) from exc
    except OSError as exc:
        raise PayrollError(f"Could not save to '{path}': {exc}") from exc
    finally:
        book.close()

    return path
