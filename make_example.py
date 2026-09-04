#!/usr/bin/env python3
"""
Generate the example workbooks that ship with this repo.

Everything here is invented. The company names are the standard fictional
placeholders (Contoso, Northwind, Globex...) precisely so they cannot be
mistaken for anyone's real client, and no person named here exists.

The synthetic data is not arbitrary: it reproduces every structural feature a
real export has, because that is what makes it useful as a fixture —
quarter-hours, an employee with no hours at all, multi-project employees,
sick/vacation/holiday charge codes, and names with parentheses, apostrophes,
accents and non-Latin scripts.

    python3 make_example.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from payroll_core import force_text, process_timesheet, save_report

HERE = Path(__file__).resolve().parent
TIMESHEET = HERE / "example_timesheet.xlsx"
REPORT = HERE / "example_report.xlsx"

PERIOD = "8/9/26 - 8/22/26"

# (employee, [(project, hours), ...])
STAFF = [
    ("Robin (Robbie) Vale", [
        ("Contoso-4100 4100", 11.5),
        ("Contoso-AXIS-1 AXIS-1", 4),
        ("Contoso-AXIS-3 AXIS-3", 7),
        ("Contoso-AXIS-4 AXIS-4", 55.5),
    ]),                                                    # four projects, 78 h
    ("Priya Raghunathan", [
        ("CD-TRAIN Trainee Program", 40),
    ]),                                                    # part time
    ("Émile Deschamps", [
        ("CD-SICK Sick Leave", 6),
        ("CD-OUTR Outreach", 61.25),
        ("CD-VAC Paid Vacation", 3),
    ]),                                                    # all three buckets, quarter hours
    ("Dana O'Neill-Barrett", [
        ("Northwind-Northwind NDA Northwind NDA", 23.25),
        ("Northwind-NW-001-018 NW-001-018", 42.75),
        ("CD-SICK Sick Leave", 7),
    ]),
    ("王小明", [
        ("Globex-Trial-A Trial-A", 27.5),
        ("Globex-Trial-B Trial-B", 27.5),
        ("Globex-Trial-C Trial-C", 23),
        ("Globex-Trial-D Trial-D", 2),
    ]),
    ("Marcus Whitfield", [
        ("CD-SICK Sick Leave", 5),
        ("Initech-INT-4180-702 INT-4180-702", 75),
    ]),
    ("Sofia Almeida", [
        ("CD-HOL Company Holiday", 8),
        ("Fabrikam-FBK-88 FBK-88", 72),
    ]),                                                    # holiday counts as work
    ("Tomasz Wiśniewski", []),                             # nobody logged any time
    ("Aisha Bello", [
        ("Aperture-Workstream One Workstream One", 33.5),
        ("Aperture-Workstream Two Workstream Two", 46.5),
    ]),
    ("Kenji Nakamura", [
        ("CD-ADMIN Office Admin", 80),
    ]),
    ("Freya Lindqvist", [
        ("Initech-INT-4180-701 INT-4180-701", 74),
        ("Initech-INT-4180-702 INT-4180-702", 2),
    ]),
    ("Hector Villalobos", [
        ("Umbra-ORBIT-2 ORBIT-2", 2),
    ]),                                                    # almost no time
    ("Nadia Petrov", [
        ("Contoso-CTX-11/12/13/14 CTX-11/12/13/14", 24),
        ("Contoso-CTX-502 CTX-502", 16),
        ("Contoso-CTX-503 CTX-503", 16),
        ("Contoso-CTX-800 CTX-800", 24),
    ]),                                                    # slashes in a project code
    ("Grace Abiodun-Clarke", [
        ("Wonka-WK-999 Other WK-999 Other", 16),
        ("Wonka-WK-US-140-9901 WK-US-140-9901 ", 32),      # trailing space, as real exports have
        ("Wonka-WK-US-141-9902  WK-US-141-9902 ", 32),     # double space, likewise
    ]),
]


def build_timesheet(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Payroll Report (jobs)"

    def row(values):
        sheet.append(values)
        for column in range(1, len(values) + 1):
            force_text(sheet.cell(row=sheet.max_row, column=column))

    row(["Payroll Report", None, None, None])
    row([PERIOD, None, None, None])
    row([None, None, None, None])
    row(["Person", "Description", "Project", "Hours"])

    for name, entries in STAFF:
        total = sum(h for _, h in entries)
        row([name, None, None, total])
        row([None, "Worked Time", None, total])
        for project, value in entries:
            row([None, None, project, value])
        row([None, None, None, None])

    book.save(path)
    return path


def main() -> int:
    build_timesheet(TIMESHEET)
    result = process_timesheet(TIMESHEET)
    save_report(result, REPORT, overwrite=True)

    print(f"Wrote {TIMESHEET.name} ({len(STAFF)} people) and {REPORT.name}")
    print(f"  period {result.period}, {result.total_hours:g} hours total")
    if result.warnings:
        print("  warnings:", result.warnings)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
